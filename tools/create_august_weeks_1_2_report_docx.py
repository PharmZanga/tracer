import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
DATA = json.loads((ROOT / "tmp" / "august-weeks-1-2-report-data.json").read_text(encoding="utf-8"))
CHARTS = ROOT / "tmp" / "august-report-charts"
OUT = ROOT / "output" / "docx" / "National_Weekly_Supply_Chain_Tracer_Report_August_Weeks_1_2_2026.docx"
GREEN = "087E4A"; DARK = "063D2C"; PALE = "EAF4EE"; INK = "173126"; ORANGE = "B65E00"; RED = "B42318"

def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), fill); tc_pr.append(shd)

def border(cell, color="C7DBD0"):
    tc_pr = cell._tc.get_or_add_tcPr(); borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders"); tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = "w:" + edge; element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag); borders.append(element)
        element.set(qn("w:val"), "single"); element.set(qn("w:sz"), "5"); element.set(qn("w:color"), color)

def set_cell_text(cell, text, bold=False, color=INK, size=8.5):
    cell.text = ""; p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
    r = p.add_run(str(text)); r.bold = bold; r.font.name = "Calibri"; r.font.size = Pt(size); r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER; border(cell)

def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.style = "Table Grid"; t.autofit = False
    for i, h in enumerate(headers):
        set_cell_text(t.rows[0].cells[i], h, True, "FFFFFF", 8); shade(t.rows[0].cells[i], DARK)
        if widths: t.rows[0].cells[i].width = Cm(widths[i])
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value, False, INK, 8)
            if widths: cells[i].width = Cm(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

def add_chart(doc, name, width=16.8):
    doc.add_picture(str(CHARTS / name), width=Cm(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

def heading(doc, text, level=1):
    p = doc.add_paragraph(); p.style = f"Heading {level}"; r=p.add_run(text); r.font.color.rgb=RGBColor.from_string(GREEN); return p

def body(doc, text, bold_prefix=None):
    p = doc.add_paragraph(); p.paragraph_format.space_after=Pt(6); p.paragraph_format.line_spacing=1.12
    if bold_prefix and text.startswith(bold_prefix):
        r=p.add_run(bold_prefix); r.bold=True; p.add_run(text[len(bold_prefix):])
    else: p.add_run(text)
    return p

def pct(v): return f"{v * 100:.1f}%"
def num(v): return f"{int(round(v)):,}"

def pooled_groups(periods, key):
    """Pool availability, MOS, and stock-status results across reporting weeks."""
    grouped = {}
    for period in periods:
        for item in period[key]:
            bucket = grouped.setdefault(item["name"], {"name": item["name"], "rows": 0, "available": 0, "quantity": 0, "amc": 0, "stockout": 0})
            bucket["rows"] += item["rows"]
            bucket["available"] += item["availability"] * item["rows"]
            bucket["quantity"] += item.get("quantity", 0)
            bucket["amc"] += item.get("amc", 0)
            bucket["stockout"] += item.get("stockout", 0)
    result = []
    for bucket in grouped.values():
        rows = bucket["rows"] or 1
        bucket["availability"] = bucket["available"] / rows
        bucket["mos"] = bucket["quantity"] / bucket["amc"] if bucket["amc"] else 0
        bucket["stockoutRate"] = bucket["stockout"] / rows
        result.append(bucket)
    return sorted(result, key=lambda item: item["availability"])

def main():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc=Document(); sec=doc.sections[0]; sec.page_width=Cm(21); sec.page_height=Cm(29.7); sec.top_margin=Cm(1.35); sec.bottom_margin=Cm(1.35); sec.left_margin=Cm(1.45); sec.right_margin=Cm(1.45)
    styles=doc.styles; styles["Normal"].font.name="Calibri"; styles["Normal"].font.size=Pt(10); styles["Normal"].font.color.rgb=RGBColor.from_string(INK)
    for name,size in (("Heading 1",16),("Heading 2",12)):
        styles[name].font.name="Calibri"; styles[name].font.size=Pt(size); styles[name].font.bold=True; styles[name].font.color.rgb=RGBColor.from_string(GREEN)
    header=sec.header.paragraphs[0]; header.text="NATIONAL WEEKLY SUPPLY CHAIN TRACER REPORT"; header.runs[0].font.size=Pt(8); header.runs[0].font.bold=True; header.runs[0].font.color.rgb=RGBColor.from_string(GREEN)
    footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER; footer.add_run("National Weekly Supply Chain Tracer Report | August 2026").font.size=Pt(8)
    # Cover
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(55); r=p.add_run("REPUBLIC OF ZAMBIA\nMINISTRY OF HEALTH"); r.bold=True; r.font.size=Pt(15); r.font.color.rgb=RGBColor.from_string(DARK)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(28); r=p.add_run("National Weekly Supply Chain\nTracer Report"); r.bold=True; r.font.size=Pt(29); r.font.color.rgb=RGBColor.from_string(DARK)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run("Weeks 1 and 2, August 2026"); r.font.size=Pt(17); r.font.color.rgb=RGBColor.from_string(GREEN)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(50); p.add_run("Reporting period: 9-16 August 2026\nPrepared by: Zanga Musakuzi, Principal Pharmacist - Data Analytics").font.size=Pt(11)
    doc.add_page_break()
    w1,w2=DATA["periods"]
    total_rows = w1["national"]["rows"] + w2["national"]["rows"]
    combined = {
        "availability": (w1["national"]["availability"] * w1["national"]["rows"] + w2["national"]["availability"] * w2["national"]["rows"]) / total_rows,
    }
    heading(doc,"A. National Overview")
    body(doc,f"The combined August Week 1 and Week 2 tracer review covers {num(w1['counts']['rows'] + w2['counts']['rows'])} submitted commodity rows from all 10 provinces. National availability declined from {pct(w1['national']['availability'])} in Week 1 to {pct(w2['national']['availability'])} in Week 2. The pooled two-week availability was {pct(combined['availability'])}.")
    table(doc,["Indicator","Week 1 - 9 Aug","Week 2 - 16 Aug"],[
        ["Availability",pct(w1['national']['availability']),pct(w2['national']['availability'])], ["Average MOS",f"{w1['national']['mos']:.2f}",f"{w2['national']['mos']:.2f}"], ["Stockout rows",num(w1['national']['stockout']),num(w2['national']['stockout'])], ["Districts submitted",f"{w1['counts']['districtSubmissionFootprint']}/116",f"{w2['counts']['districtSubmissionFootprint']}/116"],
    ],[6.8,5.4,5.4])
    add_chart(doc,"national_trend.png")
    heading(doc,"Level-of-Care Performance",2); add_chart(doc,"level_of_care.png")
    doc.add_page_break()
    heading(doc,"B. Provincial Tracer Availability")
    body(doc,"Provincial results are pooled across the two August reporting weeks. Availability is the proportion of submitted tracer commodities with stock available; MOS reflects submitted stock on hand relative to average monthly consumption.")
    add_chart(doc,"province_performance.png")
    provinces=pooled_groups([w1, w2], "provinces")
    table(doc,["Province","Availability","Avg MOS","Stockout rate"],[[p['name'].replace(' PROVINCE','').title(),pct(p['availability']),f"{p['mos']:.2f}",pct(p['stockoutRate'])] for p in provinces],[6.8,3.2,3.2,4.4])
    doc.add_page_break()
    heading(doc,"C. Stock Imbalances and Programme Performance")
    body(doc,"The stock-status profile highlights the co-existence of stockouts, understocking, adequate stock and overstock across reporting facilities. Programme results identify portfolios requiring more focused review.")
    add_chart(doc,"imbalances.png")
    add_chart(doc,"programmes.png")
    doc.add_page_break()
    heading(doc,"D. ZAMMSA Weekly Stock Position")
    central=DATA["centralReport"]; summary=central["summary"]
    body(doc,f"Matching ZAMMSA weekly inventory data was available for 7 and 14 August 2026, with the central MOS extract dated {central['label']}. The central extract listed {num(summary['listed'])} commodities; {num(summary['belowTwoMos'])} were below two MOS and {num(summary['mosDataGaps'])} had unresolved MOS data gaps.")
    add_chart(doc,"zammsa_weekly.png")
    add_chart(doc,"zammsa_central_risk.png")
    heading(doc,"E. Reporting Completeness",1)
    body(doc,f"Week 1 recorded {w1['counts']['districtSubmissionFootprint']}/116 district submissions; Week 2 recorded {w2['counts']['districtSubmissionFootprint']}/116. Missing reporting should be followed up before interpreting a district as having zero stock.")
    add_chart(doc,"reporting.png")
    doc.add_page_break()
    heading(doc,"Conclusion and Recommended Actions")
    body(doc,f"The review shows relatively stable national availability across the first two August reporting weeks, though the latest week declined to {pct(w2['national']['availability'])}. At the same time, {pct(w2['national']['stockoutRate'])} of Week 2 submitted commodity rows were stocked out, demonstrating that national availability alone should not be used as the sole operational signal.")
    table(doc,["Priority action","Recommended owner","Timing"],[
        ["Validate facilities with zero stock and high AMC; initiate redistribution where stock is available.","Provincial and district pharmacists","Immediate"],
        ["Review programmes with low availability or MOS below two months and agree replenishment actions.","Programme managers and Control Tower","Within one week"],
        ["Close missing district submissions and verify reported primary-care aggregates.","Provincial health offices","Before next weekly submission"],
        ["Use ZAMMSA category-level gaps to inform allocation and procurement follow-up.","ZAMMSA and MoH","Weekly"],
    ],[9.3,4.6,2.7])
    heading(doc,"Methodology",2); body(doc,"Availability is calculated from submitted tracer commodity rows with stock available. Months of stock is derived from stock on hand divided by average monthly consumption. Stockout and risk classifications are based on submitted quantity and MOS, while reporting completeness is assessed separately against the expected district reporting universe.")
    doc.save(OUT); print(OUT)

if __name__=="__main__": main()
