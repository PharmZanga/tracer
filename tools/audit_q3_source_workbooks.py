"""Read-only integrity audit for the June-August 2026 clean tracer workbooks."""

from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
import json
import sys

import openpyxl


ROOT = Path(r"C:\Users\Zanga Musakuzi\Desktop\NSCCU DATA ANALYSIS\PROVINCIAL  tracer SUBMISSION")
MASTER = Path(r"C:\Users\Zanga Musakuzi\Desktop\tracer dashboard\JANUARY-DECEMBER TRACER 2026 19.07.26.xlsx")
SOURCES = [
    ("2026-06-07", ROOT / "tracer summery report clean data/june/07.06.2026 SUMMARY TRACE LIST.xlsx"),
    ("2026-06-14", ROOT / "tracer summery report clean data/june/week 14.06.2026 tracer.xlsx"),
    ("2026-06-21", ROOT / "tracer summery report clean data/june/21.06.2026 tracer.xlsx"),
    ("2026-06-28", MASTER),
    ("2026-07-06", MASTER),
    ("2026-07-12", MASTER),
    ("2026-07-19", MASTER),
    ("2026-07-26", ROOT / "tracer summery report clean data/july/28.07.26 summary reports.xlsx"),
    ("2026-08-02", ROOT / "tracer summery report clean data/july/TRACER SUMMARY 02 AUAGUST 2026.xlsx"),
    ("2026-08-09", ROOT / "tracer summery report clean data/august/week 1/9.8.2026 tracer summary.xlsx"),
    ("2026-08-16", ROOT / "tracer summery report clean data/august/week 2/tracer summary 16-08-2026.xlsx"),
    ("2026-08-23", ROOT / "tracer summery report clean data/august/week 3/23.08.2026.xlsx"),
    ("2026-08-30", ROOT / "tracer summery report clean data/august/week 4/30.08.2026Tracer summary report.xlsx"),
]

FIELDS = {"PROVINCE", "DISTRICT", "FACILITY LEVEL", "FACILITY NAME", "PROGRAM", "DESCRIPTION OF ITEM", "QUANTITY", "AMC", "MOS"}


def cell_text(value):
    return str(value or "").strip().upper().replace("  ", " ")


def to_number(value):
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def date_key(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return str(value or "").strip()


def header_index(sheet):
    for row_number in range(1, min(sheet.max_row, 20) + 1):
        cells = [cell_text(sheet.cell(row_number, col).value) for col in range(1, sheet.max_column + 1)]
        cells = ["AMC" if cell in {"AVERAGE AMC", "AVG AMC"} else cell for cell in cells]
        if len(FIELDS.intersection(cells)) >= 8:
            return row_number, {name: cells.index(name) for name in FIELDS.union({"DATE"}) if name in cells}, cells
    raise ValueError(f"Could not identify tracer header in {sheet.title}")


def audit_source(report_date, path):
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    output = {"period": report_date, "source": path.name, "sheets": []}
    totals = Counter()
    province_mos = defaultdict(int)
    examples = []
    for sheet in workbook.worksheets:
        try:
            header_row, index, _ = header_index(sheet)
        except ValueError:
            continue
        if not FIELDS.issubset(index):
            continue
        seen = Counter()
        rows = 0
        for values in sheet.iter_rows(min_row=header_row + 1, values_only=True):
            if not any(value is not None for value in values):
                continue
            if "DATE" in index and date_key(values[index["DATE"]]) != report_date:
                continue
            item = values[index["DESCRIPTION OF ITEM"]]
            if not str(item or "").strip():
                continue
            rows += 1
            identity = tuple(cell_text(values[index[name]]) for name in ("PROVINCE", "DISTRICT", "FACILITY LEVEL", "FACILITY NAME", "PROGRAM", "DESCRIPTION OF ITEM"))
            seen[identity] += 1
            quantity = to_number(values[index["QUANTITY"]])
            amc = to_number(values[index["AMC"]])
            mos = to_number(values[index["MOS"]])
            if not all(identity):
                totals["missingIdentity"] += 1
            if any(value is not None and value < 0 for value in (quantity, amc, mos)):
                totals["invalidNumbers"] += 1
                if len(examples) < 20:
                    examples.append({"type": "negative value", "province": identity[0], "district": identity[1], "facility": identity[3], "item": identity[5], "quantity": quantity, "amc": amc, "mos": mos})
            if quantity and quantity > 0 and not (amc and amc > 0):
                totals["positiveQuantityWithoutAmc"] += 1
            if quantity is not None and quantity >= 0 and amc and amc > 0 and mos is not None:
                calculated = quantity / amc
                if mos >= 12 and calculated > 12:
                    totals["cappedMos"] += 1
                elif abs(calculated - mos) > 0.51:
                    totals["materialMosDifference"] += 1
                    province_mos[identity[0]] += 1
                    if len(examples) < 20:
                        examples.append({"type": "MOS mismatch", "province": identity[0], "district": identity[1], "facility": identity[3], "item": identity[5], "quantity": quantity, "amc": amc, "mos": mos, "calculatedMos": round(calculated, 2)})
        duplicate_rows = sum(count - 1 for count in seen.values() if count > 1)
        totals["rows"] += rows
        totals["duplicateIdentityRows"] += duplicate_rows
        output["sheets"].append({"sheet": sheet.title, "rows": rows, "duplicateIdentityRows": duplicate_rows})
    output["totals"] = dict(totals)
    output["mosMismatchByProvince"] = dict(sorted(province_mos.items(), key=lambda entry: entry[1], reverse=True))
    output["examples"] = examples
    return output


scope = sys.argv[1] if len(sys.argv) > 1 else ""
selected_sources = [(date, path) for date, path in SOURCES if not scope or date.startswith(scope)]
Path("tmp").mkdir(exist_ok=True)
output_path = Path(f"tmp/q3-source-workbook-audit-{scope or 'all'}.json")
results = []
for date, path in selected_sources:
    results.append(audit_source(date, path))
    output_path.write_text(json.dumps(results, indent=2), encoding="utf-8")
print(json.dumps({
    "output": str(output_path),
    "periods": [{"period": entry["period"], **entry["totals"]} for entry in results],
}, indent=2))
