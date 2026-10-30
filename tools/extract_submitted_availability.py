"""Extract reported HC/HP averages without changing stock data or source workbooks."""
from __future__ import annotations

import argparse
import gc
import hashlib
import importlib.util
import json
import math
import re
import warnings
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "src/submittedAvailability2026.js"
MONTH_DIRECTORIES = {
    "02": "February province submission", "03": "march province submission",
    "04": "april provincial submission", "05": "may provinsion submission",
    "06": "june province submission", "07": "july", "08": "august", "09": "september", "10": "october",
}


def load_generator():
    spec = importlib.util.spec_from_file_location("availability_generator", ROOT / "tools/generate_tracer_facility_data.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def numeric(value):
    if isinstance(value, bool) or value is None:
        return None
    try:
        number = float(str(value).replace(",", "").replace("%", "").strip())
    except (ValueError, TypeError):
        return None
    return number if number == number else None


def rate(value):
    number = numeric(value)
    if number is None:
        return None
    if isinstance(value, str) and "%" in value:
        number /= 100
    return number if 0 <= number <= 1 else None


def source_files(generator, period):
    explicit = {}
    for value in vars(generator).values():
        candidates = value if isinstance(value, list) else [value]
        for config in candidates:
            if isinstance(config, dict) and config.get("reportDate") == period["id"]:
                for source in config.get("rawSources", []):
                    if source["path"].suffix.lower() == ".xlsx" and source["path"].exists():
                        explicit[source["province"]] = source["path"]
    month = period["month"][5:7]
    directory = generator.RAW_SUBMISSION_ROOT / MONTH_DIRECTORIES.get(month, "__unavailable__")
    week = str(period["week"]).split()[-1]
    folders = sorted(p for p in directory.glob("*") if p.is_dir() and re.search(rf"\b{week}\b", p.name)) if directory.exists() else []
    for folder in folders:
        for path in sorted(folder.glob("*.xlsx")):
            if path.name.startswith("~$"):
                continue
            province = generator.raw_submission_province(path.name)
            if province and province not in explicit:
                explicit[province] = path
    # A manifest identifies the exact imported revision when several copies exist.
    for manifest in period.get("sourceWorkbooks", []):
        matches = list(directory.rglob(manifest["filename"])) if directory.exists() else []
        verified = [p for p in matches if hashlib.sha256(p.read_bytes()).hexdigest() == manifest["sha256"]]
        if verified:
            explicit[manifest["province"]] = verified[0]
        elif len(matches) == 1:
            # Re-saved workbooks may have different ZIP/drawing metadata. Their
            # quantity cells must independently match the published block below.
            explicit[manifest["province"]] = matches[0]
        else:
            explicit.pop(manifest["province"], None)
    return explicit


def extract(generator, path, province, period):
    rows = []
    workbook = openpyxl.load_workbook(path, read_only=False, data_only=True)
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    supported = {(f["province"], f["district"], f["facilityLevel"]) for f in period["facilities"]}
    published = {}
    data = period["commodityFacilityData"]
    dictionaries = data["dictionaries"]
    for p, d, level, _, item, _, quantity, *_ in data["rows"]:
        key = (dictionaries["provinces"][p], dictionaries["districts"][d], dictionaries["levels"][level])
        if quantity is not None and quantity > 0:
            published.setdefault(key, []).append(quantity)
    revision_verified = any(m.get("province") == province and m.get("sha256") == digest
                            for m in period.get("sourceWorkbooks", []))
    withheld = 0
    for sheet in workbook.worksheets:
        name = sheet.title.strip().upper()
        if name not in {"HC", "HP"}:
            continue
        level = "HEALTH CENTRE" if name == "HC" else "HEALTH POST"
        footer = next((r for r in range(1, sheet.max_row + 1)
                       if "PERCENTAGE AVAILABILITY" in str(sheet.cell(r, 2).value or "").upper()), None)
        if not footer:
            continue
        for header_column in range(4, sheet.max_column - 1):
            district = next((generator.canonical_district(province, sheet.cell(r, header_column).value)
                             for r in range(1, 5)
                             if generator.canonical_district(province, sheet.cell(r, header_column).value)), None)
            if not district or (province, district, level) not in supported:
                continue
            column = next((c for c in (header_column, header_column - 1, header_column + 1)
                           if any("QUANTITY" in str(sheet.cell(r, c).value or "").upper()
                                  or str(sheet.cell(r, c).value or "").strip().upper() in {"SOH", "STOCK ON HAND"}
                                  for r in range(3, 6))), header_column)
            formatted = [sheet.cell(r, column) for r in range(footer, min(footer + 4, sheet.max_row + 1))
                         if "%" in sheet.cell(r, column).number_format and rate(sheet.cell(r, column).value) is not None]
            percentage_cell = formatted[0] if formatted else sheet.cell(footer, column)
            reported = rate(percentage_cell.value)
            if reported is None:
                continue
            source_positive = [numeric(sheet.cell(r, column).value)
                               for r in range(4, footer)
                               if sheet.cell(r, 2).value and (numeric(sheet.cell(r, column).value) or 0) >= 1
                               and not any(t in str(sheet.cell(r, 2).value).upper() for t in ("AVERAGE MONTH", "PERCENTAGE AVAILABILITY"))]
            expected_positive = published.get((province, district, level), [])
            values_match = sorted(source_positive) == sorted(expected_positive)
            if not values_match and all(float(v).is_integer() for v in expected_positive):
                values_match = sorted(math.floor(v) for v in source_positive) == sorted(expected_positive)
            if not revision_verified and not values_match:
                withheld += 1
                continue
            # Header counts are read only under explicitly labelled columns.
            submitted_count = None
            for header in range(1, 4):
                for offset in (1, 2, 3):
                    label = str(sheet.cell(header, column + offset).value or "").upper()
                    if "FACILIT" in label and "REPORTED" in label:
                        count = numeric(sheet.cell(header + 1, column + offset).value)
                        if count is not None and count > 0 and count.is_integer():
                            submitted_count = int(count)
            rows.append({"province": province, "district": district, "facilityLevel": level,
                         "availability": reported, "reportingFacilities": submitted_count,
                         "workbook": path.name, "sha256": digest, "sheet": sheet.title,
                         "cell": percentage_cell.coordinate,
                         "quantityValuesVerified": values_match, "workbookRevisionVerified": revision_verified,
                         "basis": "submitted facility-average availability"})
    workbook.close()
    del workbook
    gc.collect()
    if withheld:
        print(f"{period['id']} {province}: {withheld} blocks withheld (source stock does not match)", flush=True)
    return rows


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--write", action="store_true")
    parser.add_argument("--period", action="append")
    parser.add_argument("--province", action="append")
    args = parser.parse_args()
    warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")
    openpyxl.reader.excel.find_images = lambda archive, target: ([], [])
    generator = load_generator()
    periods = []
    for filename in ("jan-feb", "mar-apr", "may-jun", "jul-aug", "sep", "oct"):
        periods.extend(json.loads((ROOT / f"public/data/tracer/2026/{filename}.json").read_text(encoding="utf-8"))["tracerReportingPeriods"])
    output = {"schema": "submitted-facility-average-v1", "periods": {}}
    if (args.period or args.province) and OUTPUT.exists():
        output = json.loads(OUTPUT.read_text(encoding="utf-8").split("export default ", 1)[1].rstrip(";\n"))
    for period in periods:
        if args.period and period["id"] not in args.period:
            continue
        records = [row for row in output["periods"].get(period["id"], []) if args.province and row["province"] not in args.province]
        for province, path in sorted(source_files(generator, period).items()):
            if args.province and province not in args.province:
                continue
            try:
                extracted = extract(generator, path, province, period)
                records.extend(extracted)
                print(f"{period['id']} {province}: {len(extracted)} submitted averages", flush=True)
            except (OSError, ValueError, KeyError) as error:
                print(f"{period['id']} {province}: unavailable ({error})", flush=True)
        unique = {(r["province"], r["district"], r["facilityLevel"]): r for r in records}
        output["periods"][period["id"]] = list(unique.values())
    if args.write:
        OUTPUT.write_text("// Generated from provincial source cells; stock quantities are not modified.\nexport default "
                          + json.dumps(output, separators=(",", ":")) + ";\n", encoding="utf-8")
    print(json.dumps({date: {"submittedAverages": len(rows), "missingFacilityWeights": sum(r["reportingFacilities"] is None for r in rows)}
                      for date, rows in output["periods"].items()}, indent=2), flush=True)


if __name__ == "__main__":
    main()
